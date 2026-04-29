#!/usr/bin/env python3
"""
WCCCE 2026 Keynote Slide Generator

Builds a PowerPoint from ``keynote.json`` (``slides[].markdown`` + ``kind`` + optional ``id`` for images).

Usage:
  python generate_slides.py
  python generate_slides.py --input keynote.json --output slides.pptx
  python generate_slides.py --images slide_images/

Image column (right, **square**): each slide uses ``slide_images/{id}.png``.
Missing files get a gray placeholder.
"""

import argparse
import json
import re
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


# ── Palette ──────────────────────────────────────────────────────────────────

BG_DARK      = RGBColor(0x1A, 0x1A, 0x2E)   # deep navy
BG_QUOTE     = RGBColor(0x16, 0x21, 0x3E)   # slightly lighter navy
ACCENT       = RGBColor(0xE9, 0x4F, 0x37)   # warm red-orange
TEXT_PRIMARY = RGBColor(0xF5, 0xF5, 0xF5)   # near-white
TEXT_DIM     = RGBColor(0xA0, 0xA8, 0xC0)   # muted blue-grey
TITLE_COLOR  = RGBColor(0xFF, 0xFF, 0xFF)


# ── Slide dimensions (16:9) ───────────────────────────────────────────────────

W = Inches(13.333)
H = Inches(7.5)

# Two-column layout: text left, square image right — matches 1024×1024 assets
FOOTER_H = Inches(0.45)
IMG_SIZE = Inches(6.05)
IMG_W = IMG_SIZE
IMG_H = IMG_SIZE
IMG_LEFT = W - Inches(0.5) - IMG_W
IMG_TOP = (H - FOOTER_H - IMG_H) / 2
TEXT_LEFT = Inches(0.65)
TEXT_WIDTH = IMG_LEFT - TEXT_LEFT - Inches(0.4)
PLACEHOLDER_FILL = RGBColor(0x2A, 0x2E, 0x48)
PLACEHOLDER_LINE = RGBColor(0x55, 0x5A, 0x78)


# ── Markdown inline parser ────────────────────────────────────────────────────

def parse_inline(text: str) -> list[tuple[str, bool, bool]]:
    """Return list of (text, bold, italic) tuples from inline markdown."""
    runs = []
    # Combined pattern: **bold**, *italic*, `code` (treated as italic here)
    pattern = re.compile(r'(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            runs.append((text[pos:m.start()], False, False))
        if m.group(2):
            runs.append((m.group(2), True, False))
        elif m.group(3):
            runs.append((m.group(3), False, True))
        elif m.group(4):
            runs.append((m.group(4), False, True))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], False, False))
    return runs or [(text, False, False)]


def add_runs(para, text: str, base_size: int, color: RGBColor,
             default_bold=False, default_italic=False):
    """Add inline-formatted runs to a paragraph."""
    for fragment, bold, italic in parse_inline(text):
        run = para.add_run()
        run.text = fragment
        run.font.size = Pt(base_size)
        run.font.color.rgb = color
        run.font.bold = bold or default_bold
        run.font.italic = italic or default_italic


# ── Low-level helpers ─────────────────────────────────────────────────────────

def blank_slide(prs: Presentation, bg: RGBColor):
    layout = prs.slide_layouts[6]   # blank
    slide = prs.slides.add_slide(layout)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return slide


def add_textbox(slide, left, top, width, height):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    return tf


def accent_bar(slide, width=Inches(0.06), height=Inches(1.2),
               left=Inches(0.55), top=Inches(1.35)):
    """Small vertical accent bar on content slides."""
    shape = slide.shapes.add_shape(
        1,  # MSO_SHAPE_TYPE.RECTANGLE
        left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()


def footer_bar(slide, text: str):
    if not text:
        return
    tf = add_textbox(slide, Inches(0.5), H - FOOTER_H + Inches(0.02), W - Inches(1), Inches(0.35))
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    run = para.add_run()
    run.text = text
    run.font.size = Pt(11)
    run.font.color.rgb = TEXT_DIM
    run.font.italic = True


def add_slide_image_column(
    slide, slide_index: int, image_dir: Path | None, slide_id: str | None = None
):
    """Right column: square image ``{slide_id}.png`` (UUID-based lookup), else gray placeholder."""
    left, top, iw, ih = IMG_LEFT, IMG_TOP, IMG_W, IMG_H
    path: Path | None = None
    if image_dir is not None and image_dir.is_dir():
        if slide_id:
            sid = slide_id.strip()
            if sid:
                for ext in (".png", ".jpg", ".jpeg", ".webp"):
                    cand = image_dir / f"{sid}{ext}"
                    if cand.is_file():
                        path = cand
                        break
    if path is not None:
        slide.shapes.add_picture(str(path), left, top, width=iw, height=ih)
        return
    ph = slide.shapes.add_shape(1, left, top, iw, ih)
    ph.fill.solid()
    ph.fill.fore_color.rgb = PLACEHOLDER_FILL
    ph.line.fill.solid()
    ph.line.fill.fore_color.rgb = PLACEHOLDER_LINE
    ph.line.width = Pt(1)
    tf = add_textbox(slide, left, top + ih / 2 - Inches(0.15), iw, Inches(0.35))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    label = (slide_id or "").strip()[:12] or f"{slide_index:02d}"
    r.text = label
    r.font.size = Pt(18)
    r.font.color.rgb = TEXT_DIM


# ── Slide builders ─────────────────────────────────────────────────────────────

def build_title_slide(
    prs,
    title: str,
    subtitle: str,
    author: str,
    footer: str,
    slide_index: int,
    image_dir: Path | None,
    slide_id: str | None = None,
):
    slide = blank_slide(prs, BG_DARK)

    col_center = TEXT_LEFT + TEXT_WIDTH / 2

    # Horizontal accent line (left column only)
    bar_w = TEXT_WIDTH - Inches(0.6)
    bar = slide.shapes.add_shape(
        1, col_center - bar_w / 2, H / 2 - Inches(0.05), bar_w, Inches(0.06)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    # Title
    tf = add_textbox(slide, TEXT_LEFT, Inches(1.05), TEXT_WIDTH, Inches(2))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    run.font.size = Pt(46)
    run.font.bold = True
    run.font.color.rgb = TITLE_COLOR

    # Subtitle
    if subtitle:
        tf2 = add_textbox(slide, TEXT_LEFT, Inches(3.15), TEXT_WIDTH, Inches(0.9))
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.size = Pt(26)
        run2.font.color.rgb = ACCENT
        run2.font.bold = False

    # Author / date
    if author:
        tf3 = add_textbox(slide, TEXT_LEFT, Inches(4.15), TEXT_WIDTH, Inches(0.6))
        p3 = tf3.paragraphs[0]
        p3.alignment = PP_ALIGN.CENTER
        run3 = p3.add_run()
        run3.text = author
        run3.font.size = Pt(18)
        run3.font.color.rgb = TEXT_DIM

    add_slide_image_column(slide, slide_index, image_dir, slide_id)
    footer_bar(slide, footer)


def build_content_slide(
    prs,
    slide_title: str,
    lines: list[str],
    footer: str,
    slide_index: int,
    image_dir: Path | None,
    slide_id: str | None = None,
):
    slide = blank_slide(prs, BG_DARK)
    accent_bar(slide, left=TEXT_LEFT - Inches(0.1))

    # Title
    tf_title = add_textbox(slide, TEXT_LEFT, Inches(0.35), TEXT_WIDTH, Inches(0.9))
    p_title = tf_title.paragraphs[0]
    p_title.alignment = PP_ALIGN.LEFT
    run_t = p_title.add_run()
    run_t.text = slide_title
    run_t.font.size = Pt(30)
    run_t.font.bold = True
    run_t.font.color.rgb = TITLE_COLOR

    # Divider line under title
    div = slide.shapes.add_shape(1, TEXT_LEFT, Inches(1.22), TEXT_WIDTH, Inches(0.025))
    div.fill.solid()
    div.fill.fore_color.rgb = RGBColor(0x40, 0x44, 0x60)
    div.line.fill.background()

    # Content area
    content_top = Inches(1.42)
    content_h = H - content_top - FOOTER_H - Inches(0.08)
    tf = add_textbox(slide, TEXT_LEFT, content_top, TEXT_WIDTH, content_h)
    tf.paragraphs[0].clear()   # remove default empty para

    first = True
    for line in lines:
        if not first:
            p = tf.add_paragraph()
        else:
            p = tf.paragraphs[0]
            first = False

        stripped = line.strip()

        if stripped.startswith('- '):
            # Bullet
            p.alignment = PP_ALIGN.LEFT
            bullet_text = stripped[2:]
            run_bullet = p.add_run()
            run_bullet.text = "  •  "
            run_bullet.font.size = Pt(22)
            run_bullet.font.color.rgb = ACCENT
            add_runs(p, bullet_text, 22, TEXT_PRIMARY)
            p.space_before = Pt(4)

        elif stripped == '':
            # Blank line → small spacer
            p.space_before = Pt(8)

        else:
            # Regular paragraph (may contain inline bold/italic)
            p.alignment = PP_ALIGN.LEFT
            add_runs(p, stripped, 20, TEXT_PRIMARY)
            p.space_before = Pt(6)

    add_slide_image_column(slide, slide_index, image_dir, slide_id)
    footer_bar(slide, footer)


def build_quote_slide(
    prs,
    lines: list[str],
    footer: str,
    slide_index: int,
    image_dir: Path | None,
    slide_id: str | None = None,
):
    """Large centered quote — no title heading."""
    slide = blank_slide(prs, BG_QUOTE)

    bar_w = TEXT_WIDTH - Inches(0.4)
    col_mid = TEXT_LEFT + TEXT_WIDTH / 2

    # Top accent bar (horizontal, left column)
    bar_top = slide.shapes.add_shape(
        1, col_mid - bar_w / 2, Inches(1.05), bar_w, Inches(0.06)
    )
    bar_top.fill.solid()
    bar_top.fill.fore_color.rgb = ACCENT
    bar_top.line.fill.background()

    # Bottom accent bar
    bar_bot = slide.shapes.add_shape(
        1, col_mid - bar_w / 2, H - Inches(1.45), bar_w, Inches(0.06)
    )
    bar_bot.fill.solid()
    bar_bot.fill.fore_color.rgb = ACCENT
    bar_bot.line.fill.background()

    # Quote text
    tf = add_textbox(slide, TEXT_LEFT, Inches(1.35), TEXT_WIDTH, H - Inches(2.35))
    tf.paragraphs[0].clear()

    first = True
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if not first:
            p = tf.add_paragraph()
            p.space_before = Pt(14)
        else:
            p = tf.paragraphs[0]
            first = False

        p.alignment = PP_ALIGN.CENTER
        add_runs(p, stripped, 24, TEXT_PRIMARY, default_italic=True)

    add_slide_image_column(slide, slide_index, image_dir, slide_id)
    footer_bar(slide, footer)


def append_markdown_slide(
    prs: Presentation,
    slide_type: str,
    content: str,
    meta: dict,
    slide_index: int,
    footer: str,
    image_dir: Path | None,
    slide_id: str | None = None,
) -> None:
    """Append one slide from a markdown block (title / quote / content)."""
    if slide_type == "title":
        lines = content.strip().splitlines()
        slide_title = ""
        extras: list[str] = []
        for line in lines:
            s = line.strip()
            if s.startswith("# ") and not slide_title:
                slide_title = s[2:]
            elif s.startswith("**") and s.endswith("**"):
                extras.append(s[2:-2])
        subtitle = extras[0] if len(extras) > 0 else meta.get("subtitle", "")
        author = extras[1] if len(extras) > 1 else meta.get("author", "")
        build_title_slide(
            prs, slide_title, subtitle, author, footer, slide_index, image_dir, slide_id
        )
        return

    if slide_type == "quote":
        all_lines = content.strip().splitlines()
        quote_lines: list[str] = []
        for line in all_lines:
            s = line.strip()
            if s.startswith("# "):
                quote_lines.append(s[2:])
            elif s:
                quote_lines.append(s)
        build_quote_slide(prs, quote_lines, footer, slide_index, image_dir, slide_id)
        return

    slide_title, body_lines = extract_title_and_body(content)
    build_content_slide(
        prs, slide_title, body_lines, footer, slide_index, image_dir, slide_id
    )


# ── Markdown helpers (deck slide bodies) ───────────────────────────────────────

def extract_title_and_body(content: str) -> tuple[str, list[str]]:
    """Pull the first # heading as title, rest as body lines."""
    lines = content.strip().splitlines()
    title = ''
    body = []
    for i, line in enumerate(lines):
        stripped = line.strip()
        if stripped.startswith('# ') and not title:
            title = stripped[2:]
        else:
            body.append(line)
    # Clean leading blank lines from body
    while body and not body[0].strip():
        body.pop(0)
    return title, body


# ── Main ───────────────────────────────────────────────────────────────────────

def generate_from_deck(
    deck_path: Path, output_path: Path, image_dir: Path | None = None
) -> None:
    """Build PPTX from ``keynote.json`` (``slides[].markdown`` + ``kind`` + stable ``id`` for images)."""
    deck = json.loads(deck_path.read_text(encoding="utf-8"))
    fm = deck.get("frontmatter") or {}
    footer = str(fm.get("footer", ""))
    slides = deck.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("Deck has no slides[]")

    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    for slide_idx, spec in enumerate(slides):
        if not isinstance(spec, dict):
            continue
        md = str(spec.get("markdown", "")).strip()
        if not md:
            continue
        kind = str(spec.get("kind", "content")).strip().lower()
        if kind not in ("title", "quote", "content"):
            kind = "content"
        slide_type = "title" if kind == "title" else ("quote" if kind == "quote" else "single")
        sid = str(spec.get("id", "")).strip() or None
        append_markdown_slide(
            prs, slide_type, md, fm, slide_idx, footer, image_dir, sid
        )

    prs.save(output_path)
    print(f"Saved: {output_path}  ({len(prs.slides)} slides) from deck")


def main():
    parser = argparse.ArgumentParser(
        description="Generate WCCCE keynote slides from keynote.json."
    )
    parser.add_argument(
        "--input",
        default="keynote.json",
        help="Input deck (keynote.json only)",
    )
    parser.add_argument("--output", default="slides.pptx", help="Output PPTX file")
    parser.add_argument(
        "--images",
        default=None,
        metavar="DIR",
        help="Square images: {slide id}.png (UUID-based lookup)",
    )
    args = parser.parse_args()

    input_path = Path(args.input)
    output_path = Path(args.output)
    image_dir = Path(args.images).resolve() if args.images else None

    if input_path.suffix.lower() != ".json":
        print(
            f"Error: --input must be a .json deck (got {input_path}).",
            file=sys.stderr,
        )
        return

    if not input_path.exists():
        print(f"Error: {input_path} not found", file=sys.stderr)
        return
    if image_dir is not None and not image_dir.is_dir():
        print(f"Error: image directory not found: {image_dir}", file=sys.stderr)
        return

    try:
        generate_from_deck(input_path, output_path, image_dir=image_dir)
    except (json.JSONDecodeError, ValueError) as e:
        print(f"Error: {e}", file=sys.stderr)


if __name__ == '__main__':
    main()
